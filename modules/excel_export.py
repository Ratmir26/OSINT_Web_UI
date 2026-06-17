import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment


def _auto_width(ws):
    for col_cells in ws.columns:
        max_len = 0
        col_letter = col_cells[0].column_letter
        for cell in col_cells:
            try:
                val = str(cell.value or "")
                max_len = max(max_len, len(val))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 3, 80)


def export_to_excel(data: dict, query: str, output_path: str) -> str:
    wb = Workbook()

    now = datetime.now().strftime("%d.%m.%Y %H:%M")

    ws1 = wb.active
    ws1.title = "Обзор"
    ws1.cell(row=1, column=1, value=f"OSINT Досье: {query}").font = Font(bold=True, size=14)
    ws1.cell(row=2, column=1, value=f"Дата: {now}")
    ws1.merge_cells("A1:B1")
    ws1.merge_cells("A2:B2")

    row = 5
    ws1.cell(row=row, column=1, value="Ключевые данные").font = Font(bold=True, size=12)

    gis = data.get("2gis", {})
    if gis:
        row += 1
        ws1.cell(row=row, column=1, value="2GIS:").font = Font(bold=True)
        row += 1
        ws1.cell(row=row, column=1, value="Название")
        ws1.cell(row=row, column=2, value=gis.get("name", "—"))
        row += 1
        ws1.cell(row=row, column=1, value="Адрес")
        ws1.cell(row=row, column=2, value=gis.get("address", "—"))
        row += 1
        ws1.cell(row=row, column=1, value="Телефоны")
        ws1.cell(row=row, column=2, value=", ".join(gis.get("phones", [])) if gis.get("phones") else "—")

    yndx = data.get("yandex", {})
    if yndx:
        row += 1
        ws1.cell(row=row, column=1, value="Яндекс.Карты:").font = Font(bold=True)
        row += 1
        ws1.cell(row=row, column=1, value="Название")
        ws1.cell(row=row, column=2, value=yndx.get("name", "—"))
        row += 1
        ws1.cell(row=row, column=1, value="Адрес")
        ws1.cell(row=row, column=2, value=yndx.get("address", "—"))
        row += 1
        ws1.cell(row=row, column=1, value="Телефоны")
        ws1.cell(row=row, column=2, value=", ".join(yndx.get("phones", [])) if yndx.get("phones") else "—")

    soc = data.get("socials", {})
    if soc:
        row += 1
        ws1.cell(row=row, column=1, value="Соцсети:").font = Font(bold=True)
        for name, url in soc.items():
            row += 1
            ws1.cell(row=row, column=1, value=name)
            ws1.cell(row=row, column=2, value=url)

    web = data.get("webpage", {})
    if web:
        row += 1
        ws1.cell(row=row, column=1, value="Контакты:").font = Font(bold=True)
        phones = web.get("phones", [])
        if phones:
            row += 1
            ws1.cell(row=row, column=1, value="Телефоны")
            ws1.cell(row=row, column=2, value=", ".join(phones))
        emails = web.get("emails", [])
        if emails:
            row += 1
            ws1.cell(row=row, column=1, value="Email")
            ws1.cell(row=row, column=2, value=", ".join(emails))

    _auto_width(ws1)

    ws2 = wb.create_sheet("Веб-результаты")
    ws2.cell(row=1, column=1, value="Title").font = Font(bold=True)
    ws2.cell(row=1, column=2, value="Link").font = Font(bold=True)
    ws2.cell(row=1, column=3, value="Snippet").font = Font(bold=True)
    for i, r in enumerate(data.get("web", []), start=2):
        ws2.cell(row=i, column=1, value=r.get("title", ""))
        ws2.cell(row=i, column=2, value=r.get("link", ""))
        ws2.cell(row=i, column=3, value=r.get("snippet", ""))
    _auto_width(ws2)

    if gis:
        ws3 = wb.create_sheet("2GIS")
        ws3.cell(row=1, column=1, value="Field").font = Font(bold=True)
        ws3.cell(row=1, column=2, value="Value").font = Font(bold=True)
        fields = [
            ("Название", gis.get("name", "—")),
            ("Адрес", gis.get("address", "—")),
            ("Рейтинг", gis.get("rating", "—")),
            ("Кол-во отзывов", str(gis.get("reviews_count", ""))),
            ("Телефоны", ", ".join(gis.get("phones", [])) if gis.get("phones") else "—"),
            ("Сайт", gis.get("website", "—")),
        ]
        if gis.get("emails"):
            fields.append(("Email", ", ".join(gis["emails"])))
        for i, (label, value) in enumerate(fields, start=2):
            ws3.cell(row=i, column=1, value=label)
            ws3.cell(row=i, column=2, value=value)
        _auto_width(ws3)

    if yndx:
        ws4 = wb.create_sheet("Яндекс.Карты")
        ws4.cell(row=1, column=1, value="Field").font = Font(bold=True)
        ws4.cell(row=1, column=2, value="Value").font = Font(bold=True)
        fields = [
            ("Название", yndx.get("name", "—")),
            ("Адрес", yndx.get("address", "—")),
            ("Телефоны", ", ".join(yndx.get("phones", [])) if yndx.get("phones") else "—"),
            ("Часы работы", yndx.get("hours", "—")),
            ("Рейтинг", yndx.get("rating", "—")),
            ("Ссылка", yndx.get("url", "—")),
        ]
        for i, (label, value) in enumerate(fields, start=2):
            ws4.cell(row=i, column=1, value=label)
            ws4.cell(row=i, column=2, value=value)
        _auto_width(ws4)

    if soc:
        ws5 = wb.create_sheet("Соцсети")
        ws5.cell(row=1, column=1, value="Platform").font = Font(bold=True)
        ws5.cell(row=1, column=2, value="URL").font = Font(bold=True)
        for i, (name, url) in enumerate(soc.items(), start=2):
            ws5.cell(row=i, column=1, value=name)
            ws5.cell(row=i, column=2, value=url)
        _auto_width(ws5)

    contacts_data = {}
    if web:
        if web.get("phones"):
            contacts_data["Телефон"] = web["phones"]
        if web.get("emails"):
            contacts_data["Email"] = web["emails"]
    if gis and gis.get("phones"):
        existing = contacts_data.get("Телефон", [])
        for p in gis["phones"]:
            if p not in existing:
                existing.append(p)
        contacts_data["Телефон"] = existing
    if yndx and yndx.get("phones"):
        existing = contacts_data.get("Телефон", [])
        for p in yndx["phones"]:
            if p not in existing:
                existing.append(p)
        contacts_data["Телефон"] = existing

    if contacts_data:
        ws6 = wb.create_sheet("Контакты")
        ws6.cell(row=1, column=1, value="Type").font = Font(bold=True)
        ws6.cell(row=1, column=2, value="Value").font = Font(bold=True)
        row = 2
        for ctype, values in contacts_data.items():
            for val in values:
                ws6.cell(row=row, column=1, value=ctype)
                ws6.cell(row=row, column=2, value=val)
                row += 1
        _auto_width(ws6)

    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    wb.save(output_path)
    return output_path


def export_to_excel_from_query(query: str, output_path: str) -> str:
    from osint_agent import run_osint
    data = run_osint(query, return_data=True)
    return export_to_excel(data, query, output_path)
